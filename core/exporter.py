"""Export a note as TXT, DOCX, XML, JSON or Markdown (all open formats)."""
import io
import json
import xml.etree.ElementTree as ET
from xml.dom import minidom


def to_txt(n):  return (n["title"] + "\n\n" + n["text"]).encode("utf-8")
def to_md(n):   return ("# %s\n\n_tags: %s_\n\n%s\n" % (n["title"], n["tags"], n["text"])).encode("utf-8")
def to_json(n): return json.dumps(n, indent=2, ensure_ascii=False).encode("utf-8")


def to_xml(n):
    root = ET.Element("note", id=str(n["id"]), created=n["created"], engine=n["engine"] or "")
    ET.SubElement(root, "title").text = n["title"]
    ET.SubElement(root, "tags").text = n["tags"]
    body = ET.SubElement(root, "content")
    for line in n["text"].splitlines():
        ET.SubElement(body, "line").text = line
    return minidom.parseString(ET.tostring(root)).toprettyxml(indent="  ", encoding="utf-8")


def to_docx(n):
    from docx import Document
    d = Document()
    d.add_heading(n["title"], 1)
    for line in n["text"].splitlines():
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


FORMATS = {
    "txt":  (to_txt,  "text/plain"),
    "md":   (to_md,   "text/markdown"),
    "json": (to_json, "application/json"),
    "xml":  (to_xml,  "application/xml"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
}
